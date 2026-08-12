from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "submission" / "budgetloop-loop-engineering"
ASSETS = OUT / "assets"
TARGET = OUT / "BudgetLoop_LoopEngineering_项目方案.docx"

FONT = "PingFang SC"
MONO = "Menlo"
BLUE = "2F6BEB"
NAVY = "0B1F3A"
MUTED = "66758A"
PALE = "F4F7FC"
LINE = "DCE5F2"
GREEN = "18A66A"
GREEN_PALE = "F0FBF6"
GOLD = "E6A817"
GOLD_PALE = "FFF8E8"
RED = "C94F4F"
WHITE = "FFFFFF"


def rgb(value):
    return RGBColor.from_string(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, bold=False, color=NAVY, font=FONT, italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def set_para(p, after=8, before=0, line=1.333, align=None, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    if align is not None:
        p.alignment = align
    return p


def add_text(doc, text, size=11, bold=False, color=NAVY, after=8, before=0,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.333, italic=False, keep=False):
    p = doc.add_paragraph()
    set_para(p, after, before, line, align, keep)
    set_run(p.add_run(text), size, bold, color, italic=italic)
    return p


def add_rich(doc, parts, after=8, before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.333):
    p = doc.add_paragraph()
    set_para(p, after, before, line, align)
    for text, opts in parts:
        set_run(p.add_run(text), **opts)
    return p


def add_title(doc, title, kicker=None, subtitle=None):
    if kicker:
        add_text(doc, kicker, 9, True, BLUE, after=4, align=WD_ALIGN_PARAGRAPH.LEFT, line=1)
    p = doc.add_paragraph()
    set_para(p, after=5, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)
    set_run(p.add_run(title), 22, True, NAVY)
    if subtitle:
        add_text(doc, subtitle, 10.5, False, MUTED, after=12, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.15)
    else:
        add_text(doc, "", 1, after=6, line=1)


def add_section_head(doc, text):
    p = doc.add_paragraph()
    set_para(p, after=5, before=7, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True)
    set_run(p.add_run(text), 13, True, BLUE)
    return p


def add_callout(doc, label, text, fill=PALE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, accent, "10")
    p = cell.paragraphs[0]
    set_para(p, after=3, line=1.15)
    set_run(p.add_run(label + "  "), 10, True, accent)
    set_run(p.add_run(text), 10.5, False, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_fact_grid(doc, facts, fills=None):
    widths = [2340] * len(facts)
    table = doc.add_table(rows=1, cols=len(facts))
    set_table_geometry(table, widths)
    for i, (label, value, note) in enumerate(facts):
        cell = table.cell(0, i)
        set_cell_shading(cell, fills[i] if fills else PALE)
        set_cell_border(cell, LINE)
        p = cell.paragraphs[0]
        set_para(p, after=3, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.05)
        set_run(p.add_run(label), 8.5, True, MUTED)
        p = cell.add_paragraph()
        set_para(p, after=2, line=1)
        set_run(p.add_run(value), 16, True, BLUE if i != len(facts) - 1 else GREEN)
        p = cell.add_paragraph()
        set_para(p, after=0, line=1.05)
        set_run(p.add_run(note), 8.3, False, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_matrix(doc, headers, rows, widths, header_fill=BLUE, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, WHITE)
        p = cell.paragraphs[0]
        set_para(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.05)
        set_run(p.add_run(h), 9, True, WHITE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], WHITE if r_idx % 2 == 0 else PALE)
            set_cell_border(cells[i], LINE)
            p = cells[i].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(value) < 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0, align=align, line=1.15)
            set_run(p.add_run(value), font_size, i == 0, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_steps(doc, steps):
    for idx, (title, body) in enumerate(steps, 1):
        p = doc.add_paragraph()
        set_para(p, after=5, line=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_run(p.add_run(f"{idx:02d}"), 10, True, WHITE)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), BLUE)
        pPr.append(shd)
        set_run(p.add_run(f"  {title}  "), 10.5, True, WHITE)
        set_run(p.add_run(body), 10, False, WHITE)


def add_image(doc, path, width=6.2, caption=None, concept=False):
    p = doc.add_paragraph()
    set_para(p, after=3, align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        label = "概念效果｜目标态  " if concept else "图示  "
        add_text(doc, label + caption, 8.5, concept, GOLD if concept else MUTED,
                 after=7, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.1)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_border(cell, NAVY)
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.12, align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, line in enumerate(lines):
        run = p.add_run(line + ("\n" if i < len(lines) - 1 else ""))
        set_run(run, 8.4, False, WHITE, MONO)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def field(paragraph, code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(NAVY)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para(hp, after=0, line=1)
    set_run(hp.add_run("BudgetLoop｜闭环工程实践挑战"), 8.5, True, MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para(fp, after=0, line=1)
    set_run(fp.add_run("BudgetLoop 项目组  ·  "), 8, False, MUTED)
    field(fp, "PAGE")


def build():
    doc = Document()
    configure_document(doc)

    # 1 封面
    add_text(doc, "闭环工程实践挑战｜项目方案", 10, True, BLUE, after=86,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    add_text(doc, "BudgetLoop", 34, True, NAVY, after=8,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    add_text(doc, "让智能体团队在预算内持续完成任务", 20, True, BLUE, after=16,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1.15)
    add_text(doc, "从自然语言目标、团队分工与工具执行，到反馈修正、测试验收与证据回放的完整闭环系统", 12.5,
             False, MUTED, after=40, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.35)
    add_callout(doc, "核心主张", "不是再造一个会回答问题的模型，而是为多个智能体执行引擎提供可预算、可审计、可介入、可复现的工程控制平面。", GREEN_PALE, GREEN)
    add_text(doc, "参赛团队：BudgetLoop 项目组", 10.5, True, NAVY, after=4, before=40,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    add_text(doc, "代码仓库：github.com/bei666qi-pan/BudgetLoop", 9.5, False, MUTED, after=3,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    add_text(doc, "提交日期：2026 年 8 月", 9.5, False, MUTED, after=0,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1)
    page_break(doc)

    # 2 执行摘要
    add_title(doc, "执行摘要", "01｜项目价值", "从“生成一次”升级为“持续完成”，再加上一层硬预算治理。")
    add_rich(doc, [
        ("BudgetLoop 是面向真实软件交付的预算感知智能体团队控制面。", {"size": 11.5, "bold": True, "color": NAVY}),
        ("系统接收明确目标与验收条件，自动拆解角色和阶段，调用代码、浏览器、测试与模型工具，并将每轮退出码、差异、测试和预算快照重新送回决策层，直到达成验收或形成可解释的部分完成报告。", {"size": 11, "bold": False, "color": NAVY}),
    ], after=12)
    add_fact_grid(doc, [
        ("真实后端验证", "960 项通过", "覆盖率 85.84%"),
        ("真实前端验证", "448 项通过", "23 个测试文件"),
        ("完整业务验收", "7 / 7", "浏览器场景通过"),
        ("闭环系统能力", "全链路", "预算、审批、回放"),
    ])
    add_section_head(doc, "为什么有竞争力")
    add_matrix(doc, ["能力", "传统单轮智能体", "BudgetLoop"], [
        ("预算", "执行后统计", "调用前原子预留，超限不请求上游"),
        ("协作", "多窗口松散拼接", "显式角色、共享上下文、任务移交与团队预算"),
        ("反馈", "依赖模型自评", "测试、编译、差异和重复动作等确定性信号"),
        ("失败", "重试或停住", "换假设、回滚、最小修复与部分完成"),
    ], [1500, 2850, 5010])
    add_callout(doc, "一句话定位", "BudgetLoop 是多个智能体执行引擎上方的“任务交付操作系统”：模型负责推理，控制面负责边界、证据和收敛。", PALE, BLUE)
    page_break(doc)

    # 3 场景与控制舱概念
    add_title(doc, "业务痛点与机会窗口", "02｜真实场景", "当智能体进入生产任务，核心问题从“会不会生成”变成“能否持续、可控地完成”。")
    add_matrix(doc, ["生产痛点", "风险", "BudgetLoop 回应"], [
        ("结果不可验证", "代码看似完整但无法运行", "用真实工具退出码和验收测试闭环"),
        ("消耗不可控", "长上下文、重复调用、无限重试", "时间、令牌、调用和费用四维硬边界"),
        ("多人协作失序", "上下文污染、责任不清", "角色、私有上下文、消息确认与任务移交"),
        ("失败不可解释", "停在半途，没有交接材料", "事件回放、策略原因和部分完成报告"),
    ], [1900, 2600, 4860])
    add_image(doc, ASSETS / "concept/agent-team-command-center-concept.png", 6.1,
              "智能体团队控制舱：预算、角色、工具、证据与验收在同一闭环内协同。", True)
    add_callout(doc, "目标体验", "评审者能够一眼看见：谁在做、做到了哪一步、还剩多少资源、为什么继续，以及什么证据证明任务已经完成。", GOLD_PALE, GOLD)
    page_break(doc)

    # 4 场景定义
    add_title(doc, "完整业务场景：生成个人预算追踪器", "03｜端到端演示", "用一个可泛化的小型应用交付任务，展示目标拆解、工具执行、反馈修正和自动验收。")
    add_section_head(doc, "用户目标")
    add_callout(doc, "输入", "生成可在浏览器直接运行的个人预算追踪器，支持收入与支出记录、余额汇总、删除、刷新后持久化、表单校验和移动端适配。", PALE, BLUE)
    add_steps(doc, [
        ("定义交付", "自然语言目标同时携带验收条件、项目路径、执行引擎和预算上限。"),
        ("组建团队", "技术统筹、界面实现、业务逻辑、审查验证四类角色分工。"),
        ("调用工具", "智能体写入网页文件，运行本地服务与浏览器自动化。"),
        ("读取反馈", "退出码、页面状态、断言结果与移动端尺寸回流控制面。"),
        ("收敛交付", "全部场景通过后输出报告；失败则进入最小修复、回滚或换假设。"),
    ])
    add_section_head(doc, "可迁移性")
    add_text(doc, "同一套闭环并不依赖预算追踪器这一固定样例。将目标与验收条件替换后，可迁移到管理后台、营销落地页、数据清洗、接口联调、测试修复和报表生成等同类任务。", 10.5)
    add_fact_grid(doc, [
        ("真实输入", "自然语言目标", "非固定提示词"),
        ("真实动作", "文件与浏览器", "工具可追踪"),
        ("真实反馈", "7 项断言", "逐项可复现"),
        ("真实结果", "可运行网页", "桌面与移动端"),
    ])
    page_break(doc)

    # 5 闭环
    add_title(doc, "闭环工程五阶段", "04｜核心方法", "每一轮都由真实执行结果驱动下一次决策，而不是让模型凭感觉宣布成功。")
    add_image(doc, ASSETS / "diagrams/loop-engineering-cycle.png", 6.3,
              "感知、决策、执行、反馈、优化五阶段持续收敛。")
    add_matrix(doc, ["阶段", "系统动作", "形成的证据"], [
        ("感知", "读取调用、工具、测试、差异和时间线", "观测记录、退出码、动作指纹"),
        ("决策", "评估预算、进展、风险和优先级", "压力模式、策略原因、审批节点"),
        ("执行", "驱动智能体与代码、浏览器等工具", "文件变化、工具输出、模型调用"),
        ("反馈", "解析验收、回归、重复动作和进展", "测试结果、进展分、预算快照"),
        ("优化", "重试、回滚、换假设或最小修复", "下一轮计划或部分完成报告"),
    ], [1200, 4450, 3710], font_size=8.8)
    page_break(doc)

    # 6 架构
    add_title(doc, "系统架构：控制权与执行权分离", "05｜技术架构", "控制平面掌握预算、状态、审批与证据；执行引擎负责推理和动作。")
    add_image(doc, ASSETS / "diagrams/system-architecture.png", 6.2,
              "五层架构把界面、控制、事实、执行和隔离证据串成一条可追溯链路。")
    add_callout(doc, "架构原则", "智能体可以被替换，但任务状态、预算、审批、工作区和最终报告始终由 BudgetLoop 控制；失败时默认关闭授权边界，而不是静默放宽。", GREEN_PALE, GREEN)
    add_text(doc, "前端通过网页接口与事件流呈现实时进展；控制面和工作进程以关系数据库作为唯一事实源，以队列承载异步执行与心跳；每次运行使用独立工作区或代码工作树，并把测试、差异和工件统一归档。", 10.5)
    page_break(doc)

    # 7 预算
    add_title(doc, "预算不是仪表盘，而是运行时控制权", "06｜原子预算", "每次模型调用前先预留资源，超限时请求不会到达上游。")
    add_fact_grid(doc, [
        ("时间", "双口径", "绝对截止 + 活跃运行"),
        ("令牌", "输入与输出", "含预留和缓存观测"),
        ("调用", "硬上限", "并发竞争同一预算行"),
        ("费用", "可配置", "未定价时不做结论"),
    ])
    add_section_head(doc, "调用前原子预留")
    add_text(doc, "预算管理器使用带条件的数据库更新，同时检查已用量、预留量、调用上限、令牌上限、费用上限和截止时间。并发调用竞争同一行，只有满足所有约束的事务能够成功；零行更新即代表拒绝执行。", 10.7)
    add_code(doc, [
        "开始调用 → 读取预算快照",
        "原子预留 → 检查时间、令牌、调用与费用",
        "预留成功 → 调用执行引擎 → 结算真实用量",
        "预留失败 → 切换策略或输出部分完成报告",
    ])
    add_section_head(doc, "压力模式")
    add_matrix(doc, ["模式", "触发信号", "动作倾向"], [
        ("常规", "预算充足且进展稳定", "完整规划、并行验证"),
        ("保守", "剩余资源收紧或连续低分", "缩短上下文、减少探索、优先关键路径"),
        ("临界", "接近硬上限或截止时间", "最小修复、冻结扩展、准备部分完成"),
    ], [1500, 3500, 4360])
    add_callout(doc, "可信口径", "真实运行的费用字段未配置价格，因此本方案不把零费用展示宣传为成本节省。", GOLD_PALE, GOLD)
    page_break(doc)

    # 8 团队
    add_title(doc, "智能体团队：显式角色、上下文与任务移交", "07｜多角色协作", "不是多个聊天窗口的拼接，而是一套有责任边界和过程记录的协作协议。")
    add_matrix(doc, ["角色", "核心职责", "阶段门禁"], [
        ("技术统筹", "拆解目标、定义接口、分配团队预算", "确认可执行计划与验收条件"),
        ("界面实现", "页面结构、视觉样式、响应式适配", "桌面与移动端关键视图可用"),
        ("业务逻辑", "交易增删、统计、持久化与校验", "数据状态和边界行为正确"),
        ("审查验证", "运行测试、解释失败、提出最小修复", "七项浏览器场景全部通过"),
    ], [1700, 4210, 3450])
    add_section_head(doc, "协作机制")
    add_steps(doc, [
        ("共享上下文", "项目目标、公共决策与跨角色证据对团队可见。"),
        ("私有上下文", "每个工作会话保留角色相关信息，降低无关上下文污染。"),
        ("消息确认", "消息具备发送、投递、确认与失败状态，可审计未读或阻塞。"),
        ("任务移交", "交付物、未决问题、预算和验收状态随移交一起持久化。"),
        ("人工审批", "关键风险点可暂停、确认、拒绝并触发重新规划。"),
    ])
    add_callout(doc, "团队预算", "总预算可按阶段与角色分配；控制面同时防止单个角色耗尽全局资源，让协作在共享目标下保持资源纪律。", GREEN_PALE, GREEN)
    page_break(doc)

    # 9 时序
    add_title(doc, "从目标到证据：每一步都可回放", "08｜协作时序", "任务、预算与验收证据沿同一事件链流动，避免“做了什么却无法证明”。")
    add_image(doc, ASSETS / "diagrams/agent-team-budget-sequence.png", 6.25,
              "智能体团队与预算控制时序：预算检查发生在动作之前，验收证据回到控制平面。")
    add_callout(doc, "闭环关键点", "验证工具不是最后才运行的附件，而是每一轮决策的事实来源；测试失败会改变下一步策略，测试通过才允许进入交付终态。", PALE, BLUE)
    page_break(doc)

    # 10 执行
    add_title(doc, "执行层：工具调用、隔离与人工确认", "09｜可运行系统", "把模型输出约束在可授权的工作区，把高风险动作置于可审批的阶段门禁。")
    add_matrix(doc, ["能力", "系统做法", "业务价值"], [
        ("执行引擎", "统一适配多类智能体编码引擎", "避免绑定单一模型或供应商"),
        ("工具事实", "将命令、文件、浏览器和测试结果归一化", "同一反馈机制覆盖多种任务"),
        ("工作区隔离", "每次运行使用独立目录或代码工作树", "降低并发污染和越权写入风险"),
        ("审批闸门", "关键风险点暂停，人工可批准或拒绝", "保留业务责任与可控介入"),
        ("失败收敛", "有界重试、换假设、回滚、最小修复", "防止无限循环和资源浪费"),
    ], [1700, 4200, 3460])
    add_section_head(doc, "关键过程记录")
    add_text(doc, "每轮计划、工具参数、退出码、标准输出、标准错误、文件差异、测试摘要、预算快照、策略切换原因、审批结果和最终报告都拥有持久化事件。前端按序号去重并支持断线回放。", 10.7)
    add_callout(doc, "故障策略", "网关不可达、工具超时、测试失败或预算不足时，系统不会假装成功：它会重试、切换策略、等待审批或产出带证据的部分完成结果。", GOLD_PALE, GOLD)
    add_section_head(doc, "泛化边界")
    add_text(doc, "BudgetLoop 控制的是“目标—动作—证据—决策”的通用协议。只要任务能够被工具执行并由可观察结果验证，就可以复用这套闭环，不局限于网页生成。", 10.5)
    page_break(doc)

    # 11 真实运行
    add_title(doc, "真实运行证据：闭环在预算内完成", "10｜实测证据", "完成态运行记录经过中文化信息重排，原始证据保留于项目工作区。")
    add_image(doc, ASSETS / "real/真实运行摘要.png", 6.3,
              "一次真实受管运行：三轮闭环、七次模型调用、总用时一分二十八秒。")
    add_callout(doc, "实测", "运行状态为已完成；调用上限为二十次、时间上限为十五分钟。该运行未配置模型价格，费用字段不用于成本结论。", GREEN_PALE, GREEN)
    add_text(doc, "这张证据证明 BudgetLoop 能把轮次、调用、时间与终态放在同一运行记录中；它不是概念界面，也没有对真实粗糙成品进行产品化重绘。", 10.5)
    page_break(doc)

    # 12 真实桌面
    add_title(doc, "真实输出：预算追踪器桌面端", "11｜业务结果", "保留智能体生成物的真实局部样貌，用可运行与可验证回应“只有概念设计”的风险。")
    add_image(doc, ASSETS / "real/真实预算追踪器桌面.png", 6.25,
              "真实生成物局部截图：标题、当前余额、总收入与总支出。")
    add_section_head(doc, "输入—输出闭环")
    add_matrix(doc, ["输入要求", "输出表现", "验收方式"], [
        ("添加收入", "收入与余额同步变化", "浏览器自动化断言数值与列表"),
        ("添加支出", "支出汇总并影响余额", "断言收入与支出同时存在"),
        ("删除交易", "列表和统计即时更新", "触发删除并检查记录消失"),
        ("刷新保留", "交易记录从本地存储恢复", "刷新页面后再次读取记录"),
    ], [2300, 3400, 3660])
    add_callout(doc, "真实边界", "本页展示的是可运行原型，不将其包装成已完成的商业产品界面；产品化目标态在下一页单独标注为概念效果。", GOLD_PALE, GOLD)
    page_break(doc)

    # 13 概念目标态
    add_title(doc, "产品化目标态：从可运行原型到财务工作台", "12｜概念效果", "在不混淆真实能力的前提下，展示 BudgetLoop 可进一步交付的高完成度体验。")
    add_image(doc, ASSETS / "concept/budget-tracker-future-concept.png", 6.25,
              "预算追踪器目标态：总览、分类支出、月度趋势、近期交易和预算进度整合在同一工作台。", True)
    add_matrix(doc, ["层级", "当前证据", "目标态扩展"], [
        ("可运行", "交易增删、汇总、持久化、校验", "账户、分类、预算与目标管理"),
        ("可验证", "七项浏览器自动化场景", "可访问性、视觉回归与多账户测试"),
        ("可运营", "本地单页应用", "云端同步、权限、审计与报表导出"),
    ], [1700, 3700, 3960])
    add_callout(doc, "说明", "本页为明确标注的概念效果，不作为已完成界面截图；它用于说明系统交付能力的产品化方向。", GOLD_PALE, GOLD)
    page_break(doc)

    # 14 移动端与验收
    add_title(doc, "七项浏览器验收：真实输入、真实输出", "13｜质量闭环", "不是只跑单个固定样例，而是覆盖创建、计算、删除、刷新、错误输入与响应式行为。")
    add_image(doc, ASSETS / "real/真实预算追踪器移动端.png", 5.45,
              "三百九十乘八百四十四验收视口：交易历史可见且无横向溢出。")
    add_matrix(doc, ["验收场景", "检查点", "结果"], [
        ("页面加载", "标题、表单与统计区域可见", "通过"),
        ("添加收入", "列表新增且余额同步", "通过"),
        ("添加支出", "收支并存且计算正确", "通过"),
        ("删除交易", "删除后列表更新", "通过"),
        ("刷新持久化", "刷新后交易仍存在", "通过"),
        ("空值校验", "无效提交不生成交易", "通过"),
        ("移动端", "无横向溢出", "通过"),
    ], [2200, 5360, 1800], font_size=8.5, header_fill=GREEN)
    page_break(doc)

    # 15 反馈异常
    add_title(doc, "反馈、修正与异常收敛", "14｜闭环可靠性", "系统关注“为什么继续”与“何时停止”，而不是把重试次数当作智能。")
    add_matrix(doc, ["反馈信号", "判断", "可选动作"], [
        ("失败测试减少", "进展为正", "继续当前假设，扩大验证"),
        ("编译错误增加", "发生回归", "回滚或进入最小修复"),
        ("动作指纹重复", "可能陷入循环", "改变假设或更换执行角色"),
        ("预算快速燃尽", "资源压力升高", "缩短上下文、冻结非关键功能"),
        ("高风险工具请求", "需要责任确认", "暂停并请求人工审批"),
        ("无法完整交付", "继续执行收益不足", "生成部分完成报告与下一步"),
    ], [2600, 2700, 4060])
    add_section_head(doc, "确定性进展评分")
    add_text(doc, "评分使用通过测试变化、失败测试变化、编译错误、差异规模、新观测数量和重复动作等公开权重的事实信号。模型不为自己的工作打分，控制面据此选择继续、回滚或换假设。", 10.7)
    add_callout(doc, "异常处理", "所有重试都有次数与预算边界；审批等待不消耗活跃运行时间，但仍受绝对截止时间约束。", PALE, BLUE)
    add_section_head(doc, "部分完成也是可信交付")
    add_text(doc, "当预算不足或外部依赖不可用时，最终报告明确列出已完成内容、未通过验收、修改文件、失败证据和建议下一步。相比静默失败，这让人可以安全接手并继续推进。", 10.5)
    page_break(doc)

    # 16 AI 工具
    add_title(doc, "人工智能工具使用说明", "15｜工具与边界", "模型承担规划、编码与解释，BudgetLoop 承担预算、权限、事实与终态判定。")
    add_matrix(doc, ["工具类别", "在场景中的作用", "进入闭环的反馈"], [
        ("模型网关", "理解目标、拆解任务、生成与修正代码", "调用状态、令牌、费用与模型标识"),
        ("智能体执行引擎", "组织推理、文件编辑和工具动作", "标准化消息、动作和工具结果"),
        ("命令与文件工具", "创建网页文件、启动服务、执行测试", "退出码、输出、差异和工件"),
        ("浏览器自动化", "模拟收入、支出、删除、刷新与移动端", "七项断言、截图和失败位置"),
        ("代码版本工具", "隔离工作区、记录修改、支持回滚", "变更集、工作树状态和提交证据"),
        ("事件流与接口", "展示实时阶段、预算、审批和报告", "带序号事件与断线回放"),
    ], [1900, 3900, 3560])
    add_section_head(doc, "人工确认节点")
    add_text(doc, "删除、外部发布、敏感路径写入或预算策略升级等关键风险动作可配置为审批闸门。人工拒绝不会丢失上下文，而会成为下一轮重新规划的反馈。", 10.7)
    add_section_head(doc, "人工智能透明度")
    add_text(doc, "界面区分计划、工具事实、模型建议和系统终态；评审者能够追溯一个结论究竟来自模型推理，还是来自实际测试与工具输出。", 10.7)
    add_callout(doc, "核心分工", "模型可以提出“我认为已经完成”，但只有控制面读取到验收证据后才能把任务标记为完成。", GREEN_PALE, GREEN)
    page_break(doc)

    # 17 真实测试
    add_title(doc, "真实验证报告", "16｜测试证据", "三层测试均在当前工作区重新执行；不降低门槛、不跳过失败、不用模拟输出替代。")
    add_fact_grid(doc, [
        ("后端", "960 通过", "1 跳过 · 53.72 秒"),
        ("覆盖率", "85.84%", "门槛 70%"),
        ("网页端", "448 通过", "23 文件 · 7.06 秒"),
        ("业务验收", "7 / 7", "11.9 秒"),
    ])
    add_image(doc, ASSETS / "real/真实令牌观测.png", 6.05,
              "真实运行令牌观测：总令牌约八点六八万，包含输入、输出与缓存统计。")
    add_section_head(doc, "测试环境隔离修复")
    add_text(doc, "本轮仅调整测试基础设施：显式隔离模型网关配置；每个数据库测试重建结构；接口内部提交运行于嵌套事务中并由外层清理。未修改公共接口、数据库迁移或生产业务契约。", 10)
    add_callout(doc, "跳过说明", "唯一跳过项位于真实团队观测端到端路径，需要额外运行时条件；它被明确记录为跳过，不伪装为通过。", GOLD_PALE, GOLD)
    page_break(doc)

    # 18 效能
    add_title(doc, "效能与质量提升说明", "17｜业务价值", "用“实测数据、已具备能力、目标效果”三层标签区分证据强度。")
    add_matrix(doc, ["标签", "结论", "证据口径"], [
        ("实测数据", "后端、网页端与七项业务验收全部达到预期", "本轮真实命令输出与截图"),
        ("已具备能力", "预算拦截、压力模式、任务移交、审批、重试与事件回放", "代码、接口、测试和规范可定位"),
        ("目标效果", "交付周期、人工步骤、无效调用和令牌消耗进一步下降", "待同任务对照评测，不冒充已完成实验"),
    ], [1800, 4550, 3010])
    add_section_head(doc, "目标效果｜待对照评测")
    add_fact_grid(doc, [
        ("小型应用交付", "30–90 分钟", "目标：替代 6–12 人工小时"),
        ("人工节点", "3 类", "目标确认、风险审批、最终验收"),
        ("令牌节省", "30%–50%", "相对无预算基线目标"),
        ("无效调用", "下降 30%+", "重复与低收益动作目标"),
    ], [PALE, PALE, GOLD_PALE, GOLD_PALE])
    add_section_head(doc, "测量方法")
    add_text(doc, "仓库提供评估脚本，可在相同任务、相同模型和相同验收条件下比较无预算、固定预算与动态预算三种策略，记录完成率、总用时、调用数、令牌、人工介入和回归次数。", 10.7)
    add_callout(doc, "不夸大事实", "上述幅度是产品化目标，不写成已完成的随机对照或大规模 A/B 实验；真实数据始终单独列示。", GOLD_PALE, GOLD)
    page_break(doc)

    # 19 运行说明
    add_title(doc, "本地运行与无视频演示", "18｜交付方式", "参赛包提供可运行代码、本地启动说明、八分钟现场脚本与离线备用证据。")
    add_section_head(doc, "依赖环境")
    add_text(doc, "Docker 与 Docker Compose v2、Git；如需真实模型执行，配置合法授权的兼容模型网关地址、密钥和模型别名。", 10.5)
    add_section_head(doc, "一键启动")
    add_code(doc, [
        "git clone https://github.com/bei666qi-pan/BudgetLoop.git",
        "cd BudgetLoop",
        "cp .env.example .env",
        "docker compose up -d --build",
        "浏览器访问 http://localhost:3000",
    ])
    add_section_head(doc, "无模型费用的离线复核")
    add_code(doc, [
        "cd web && npm test",
        "npx playwright test e2e/verify-budget-tracker.spec.ts --reporter=line",
        "预期：二十三个测试文件、四百四十八项测试、七项浏览器场景全部通过",
    ])
    add_steps(doc, [
        ("两分钟定位", "说明传统智能体缺少预算、证据和失败收敛。"),
        ("两分钟闭环", "展示任务拆解、团队分工、预算预留与策略切换。"),
        ("三分钟证据", "展示真实运行、真实网页输出和七项浏览器验收。"),
        ("一分钟价值", "总结控制平面、可替换引擎与产品化路线。"),
    ])
    add_callout(doc, "无视频提交", "本项目按要求不制作视频或幻灯片；正式方案、效果示意图、真实截图、演示脚本和可运行代码共同构成完整成果。", GREEN_PALE, GREEN)
    page_break(doc)

    # 20 核心代码与竞争
    add_title(doc, "核心代码说明与竞争优势", "19｜工程可信度", "创新集中在预算控制、确定性反馈、团队治理与可观测闭环，而非重复造轮子。")
    add_matrix(doc, ["模块", "核心位置", "作用"], [
        ("原子预算", "后端预算管理模块", "调用前预留、结算与拒绝"),
        ("压力模式", "后端策略模块", "按剩余资源切换执行强度"),
        ("进展评分", "后端评分模块", "基于测试、编译、差异和重复动作"),
        ("团队协作", "后端协作与工作容器模块", "会话、消息、任务移交与审批"),
        ("执行引擎", "后端执行引擎注册表", "统一多类编码智能体"),
        ("事件回放", "后端事件与网页事件流", "序号去重、断线续传与最终报告"),
    ], [1900, 3400, 4060], font_size=8.7)
    add_section_head(doc, "开源设计来源")
    add_text(doc, "系统借鉴成熟智能体执行内核、代码工作区权限模型、多协议模型网关和状态图编排设计，把研发资源投入到差异化控制面。设计来源在核心代码说明中逐项列出。", 10.3)
    add_section_head(doc, "竞赛级差异化")
    add_matrix(doc, ["维度", "BudgetLoop 亮点"], [
        ("完整度", "有可运行演示、代码仓库、测试、运行说明和真实输入输出"),
        ("新颖性", "把团队协作与硬预算控制组合为同一个闭环控制平面"),
        ("可信度", "完成态由真实测试判定，过程可回放，失败可解释"),
        ("扩展性", "执行引擎可替换，任务协议可迁移到多类生产流程"),
    ], [1800, 7560])
    add_callout(doc, "竞争定位", "BudgetLoop 不与基础模型竞争，而是成为企业采用多个智能体时不可缺少的预算、治理与交付层。", GREEN_PALE, GREEN)
    page_break(doc)

    # 21 路线图与提交映射
    add_title(doc, "路线图与提交完整性", "20｜从参赛作品到生产平台", "当前版本已经覆盖闭环核心要求，并为产品化治理、评测和部署预留清晰路径。")
    add_matrix(doc, ["阶段", "目标", "关键成果"], [
        ("当前", "可运行闭环与真实业务验收", "预算、团队、工具、反馈、报告、七项场景"),
        ("近期", "规模化评测与体验升级", "对照实验、可访问性、视觉回归、团队模板市场"),
        ("中期", "企业治理与多项目运营", "组织预算、权限策略、审计导出、私有化部署"),
        ("远期", "跨引擎智能体交付基础设施", "统一任务协议、策略学习与成本质量最优路由"),
    ], [1600, 3300, 4460])
    add_section_head(doc, "赛事要求逐项映射")
    add_matrix(doc, ["统一提交要求", "本包成果", "状态"], [
        ("可运行演示或本地运行说明", "代码仓库、容器启动与演示指南", "已提供"),
        ("代码仓库或核心代码说明", "公开仓库与核心代码指南", "已提供"),
        ("系统架构、依赖、用例", "主方案与参赛包说明", "已提供"),
        ("完整业务场景", "个人预算追踪器七项验收", "已提供"),
        ("效能或质量说明", "真实测试、能力证据与目标指标", "已提供"),
        ("过程可视化与关键记录", "运行摘要、令牌观测、流程图与回放", "已提供"),
        ("视频", "赛事可选；本项目明确不提交", "不适用"),
    ], [3600, 4060, 1700], font_size=8.3, header_fill=GREEN)
    add_callout(doc, "最终陈述", "BudgetLoop 让智能体从“回答问题”迈向“持续完成任务”：每个动作有预算，每次判断有证据，每个失败有去向，每次交付可复现。", GREEN_PALE, GREEN)
    add_text(doc, "BudgetLoop 项目组｜2026", 9.5, True, BLUE, after=0,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1)

    props = doc.core_properties
    props.title = "BudgetLoop 闭环工程实践挑战项目方案"
    props.subject = "智能体团队协作与预算控制"
    props.author = "BudgetLoop 项目组"
    props.keywords = "BudgetLoop, 闭环工程, 智能体团队, 预算控制"
    props.comments = "参赛提交版；概念效果与真实证据已分别标注。"
    doc.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    build()
