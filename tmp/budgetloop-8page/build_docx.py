from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/qi/Desktop/BudgetLoop")
WORK = ROOT / "tmp/budgetloop-8page"
ASSETS = WORK / "assets"
OUT = ROOT / "BudgetLoop_专业Agent团队控制面_8页方案.docx"

# Use a standalone Unicode TTF for Word/LibreOffice interoperability. TTC-based
# macOS CJK families were not resolved by LibreOffice's isolated render profile.
FONT = "Arial Unicode MS"
FONT_FILE = Path(
    "/System/Library/AssetsV2/com_apple_MobileAsset_Font8/"
    "86ba2c91f017a3749571a82f2c6d890ac7ffb2fb.asset/AssetData/PingFang.ttc"
)

# standard_business_brief preset + named BudgetLoop brand overrides.
HEADING_BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6F82"
LIGHT_MUTED = "8B98A9"
TABLE_FILL = "F2F4F7"
LINE = "D8E1EC"
WHITE = "FFFFFF"
BRAND_BLUE = "2F6BEB"  # named override: BudgetLoop brand accent
GREEN = "15956A"  # named override: verified/controlled state
GREEN_FILL = "EEF9F5"
BLUE_FILL = "EEF4FF"
GOLD = "B87800"
GOLD_FILL = "FFF8E8"
RED = "B84A4A"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, size=11, bold=False, color=INK, italic=False, font=FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def set_style_font(style, font=FONT):
    """Set every OOXML font slot so CJK text does not fall back or disappear."""
    style.font.name = font
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font)


def register_document_font(doc, font=FONT):
    """Register the CJK family in fontTable.xml for LibreOffice resolution."""
    for part in doc.part.package.parts:
        if str(part.partname) != "/word/fontTable.xml":
            continue
        root = etree.fromstring(part.blob)
        existing = root.xpath(f'./w:font[@w:name="{font}"]', namespaces={"w": root.nsmap["w"]})
        if existing:
            return
        font_el = OxmlElement("w:font")
        font_el.set(qn("w:name"), font)
        charset = OxmlElement("w:charset")
        charset.set(qn("w:val"), "86")
        family = OxmlElement("w:family")
        family.set(qn("w:val"), "swiss")
        pitch = OxmlElement("w:pitch")
        pitch.set(qn("w:val"), "variable")
        font_el.extend((charset, family, pitch))
        root.append(font_el)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return


def set_para(p, before=0, after=6, line=1.10, align=WD_ALIGN_PARAGRAPH.LEFT,
             keep=False, keep_together=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    fmt.keep_together = keep_together
    p.alignment = align
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LINE, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run(run, size=8.5, color=MUTED)


def add_text(doc, text, size=11, bold=False, color=INK, before=0, after=6,
             line=1.10, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False, keep=False):
    p = doc.add_paragraph()
    set_para(p, before, after, line, align, keep)
    set_run(p.add_run(text), size, bold, color, italic)
    return p


def add_rich(doc, parts, before=0, after=6, line=1.10,
             align=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    p = doc.add_paragraph()
    set_para(p, before, after, line, align, keep)
    for text, opts in parts:
        set_run(p.add_run(text), **opts)
    return p


def add_page_title(doc, num, title, subtitle):
    add_text(doc, f"{num:02d}  |  BUDGETLOOP", 8.5, True, BRAND_BLUE, after=4, line=1, keep=True)
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(title)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    add_text(doc, subtitle, 9.5, False, MUTED, after=10, line=1.15)


def add_heading(doc, text, level=2):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_callout(doc, label, text, fill=BLUE_FILL, accent=BRAND_BLUE, after=8):
    p = doc.add_paragraph()
    set_para(p, before=2, after=after, line=1.15, keep_together=True)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    set_run(p.add_run(label + "  "), 10, True, accent)
    set_run(p.add_run(text), 10.2, False, INK)
    return p


def add_matrix(doc, headers, rows, widths, font_size=9.3, first_col_bold=True):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, HEADING_BLUE)
        set_cell_borders(cell, WHITE, "8")
        p = cell.paragraphs[0]
        set_para(p, after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(header), 9.2, True, WHITE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_shading(cells[idx], WHITE if r_idx % 2 == 0 else TABLE_FILL)
            set_cell_borders(cells[idx])
            p = cells[idx].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0, line=1.12, align=align)
            set_run(p.add_run(value), font_size, first_col_bold and idx == 0, INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_fact_strip(doc, facts):
    table = doc.add_table(rows=1, cols=len(facts))
    widths = [9360 // len(facts)] * len(facts)
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    for idx, (label, value, note, color) in enumerate(facts):
        cell = table.cell(0, idx)
        set_cell_shading(cell, TABLE_FILL if idx % 2 == 0 else WHITE)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        set_para(p, after=2, line=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p.add_run(label), 8.3, True, MUTED)
        p2 = cell.add_paragraph()
        set_para(p2, after=2, line=1, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p2.add_run(value), 14, True, color)
        p3 = cell.add_paragraph()
        set_para(p3, after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_run(p3.add_run(note), 8.2, False, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_image(doc, path, width, alt, caption=None):
    p = doc.add_paragraph()
    set_para(p, after=3, line=1, align=WD_ALIGN_PARAGRAPH.CENTER, keep_together=True)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    drawings = run._r.xpath(".//wp:docPr")
    if drawings:
        drawings[0].set("descr", alt)
        drawings[0].set("title", alt[:80])
    if caption:
        cp = doc.add_paragraph(style="Caption")
        cp.add_run(caption)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def pil_font(size, bold=False):
    return ImageFont.truetype(str(FONT_FILE), size=size, index=2 if bold else 3)


def centered(draw, box, text, font, fill):
    x1, y1, x2, y2 = box
    bounds = draw.textbbox((0, 0), text, font=font)
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.text(((x1 + x2 - width) / 2, (y1 + y2 - height) / 2 - bounds[1]), text, font=font, fill=fill)


def rounded(draw, box, fill, outline=LINE, radius=24, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline="#" + outline, width=width)


def arrow(draw, start, end, color="#2F6BEB", width=6):
    draw.line((start, end), fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) > abs(ey - sy):
        sign = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - sign * 18, ey - 11), (ex - sign * 18, ey + 11)]
    else:
        sign = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 11, ey - sign * 18), (ex + 11, ey - sign * 18)]
    draw.polygon(pts, fill=color)


def crop_screenshots():
    crops = {
        "smart-routing.png": ("01-smart-routing-ready.png", (30, 80, 1245, 710)),
        "engine-selection.png": ("01c-engine-list.png", (30, 75, 1245, 650)),
        # Keep the real Session count, role states and Handoff header while
        # excluding run IDs, branch identifiers and workspace paths below.
        "agent-team.png": ("02-agent-team-focused.png", (20, 75, 940, 405)),
        "run-observatory.png": ("03-run-observatory-focused.png", (25, 225, 1255, 710)),
    }
    for out_name, (src_name, box) in crops.items():
        source = Image.open(ASSETS / src_name).convert("RGB")
        source.crop(box).save(ASSETS / out_name, quality=95)


def build_architecture_diagram():
    im = Image.new("RGB", (1800, 880), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 35), "目标架构：执行引擎可替换，控制权始终留在 BudgetLoop", font=pil_font(38, True), fill="#" + INK)
    d.text((62, 92), "智能选择、统一 Session、专业协作与证据收敛形成同一条可审计链路", font=pil_font(22), fill="#" + MUTED)
    layers = [
        (70, 190, 315, 340, "用户目标", "自然语言 · 文件 · 验收", BLUE_FILL, BRAND_BLUE),
        (380, 190, 690, 340, "智能路由", "任务分类 · 团队 · 引擎 · 预算", BLUE_FILL, BRAND_BLUE),
        (760, 155, 1200, 375, "BudgetLoop 控制平面", "Session 总线 · 主管 Agent · 裁判 Agent\n预算账本 · 审批 · 审计 · 终态", GREEN_FILL, GREEN),
        (1280, 190, 1715, 340, "专业 Agent 团队", "独立 Session · 私有上下文 · Skills", BLUE_FILL, BRAND_BLUE),
    ]
    for x1, y1, x2, y2, title, sub, fill, accent in layers:
        d.rounded_rectangle((x1, y1, x2, y2), radius=24, fill="#" + fill, outline="#" + accent, width=4)
        centered(d, (x1, y1 + 18, x2, y1 + 78), title, pil_font(27, True), "#" + accent)
        lines = sub.split("\n")
        for idx, line in enumerate(lines):
            centered(d, (x1 + 12, y1 + 82 + idx * 38, x2 - 12, y1 + 125 + idx * 38), line, pil_font(18), "#" + INK)
    arrow(d, (315, 265), (380, 265))
    arrow(d, (690, 265), (760, 265))
    arrow(d, (1200, 265), (1280, 265))
    engines = [("OpenHands", 85), ("Codex", 405), ("Gemini CLI", 725), ("更多适配器", 1045)]
    for label, x in engines:
        d.rounded_rectangle((x, 520, x + 255, 640), radius=20, fill="white", outline="#" + LINE, width=3)
        centered(d, (x, 535, x + 255, 590), label, pil_font(23, True), "#" + DARK_BLUE)
        centered(d, (x, 585, x + 255, 625), "受管执行引擎", pil_font(17), "#" + MUTED)
    d.rounded_rectangle((1375, 490, 1715, 675), radius=22, fill="#" + GOLD_FILL, outline="#" + GOLD, width=3)
    centered(d, (1390, 510, 1700, 565), "工具与工作区", pil_font(25, True), "#" + GOLD)
    centered(d, (1390, 570, 1700, 610), "文档 · 表格 · PPT · 代码", pil_font(18), "#" + INK)
    centered(d, (1390, 612, 1700, 650), "测试 · 浏览器 · 数据", pil_font(18), "#" + INK)
    arrow(d, (1500, 340), (1500, 490), color="#15956A")
    arrow(d, (1375, 585), (1200, 370), color="#15956A")
    d.text((1050, 740), "证据回流：工具事实 → 裁判判定 → 定向返工 / 通过 → 主管汇总", font=pil_font(21, True), fill="#" + GREEN)
    im.save(ASSETS / "architecture.png")


def build_role_diagram():
    im = Image.new("RGB", (1800, 700), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 30), "三层协作模型", font=pil_font(38, True), fill="#" + INK)
    d.text((62, 84), "主管负责全局，专家负责执行，裁判保持独立", font=pil_font(22), fill="#" + MUTED)
    boxes = [
        (70, 180, 510, 500, "主管 Agent", ["拆解目标与阶段", "分配角色与预算", "处理 Handoff", "汇总统一成果"], BLUE_FILL, BRAND_BLUE),
        (680, 180, 1120, 500, "专业 Agent", ["独立 Session", "私有上下文", "角色 Skills", "可验证交付"], GREEN_FILL, GREEN),
        (1290, 180, 1730, 500, "裁判 Agent", ["独立审查证据", "控制对话预算", "通过 / 返工 / 阻塞", "定向触发复核"], GOLD_FILL, GOLD),
    ]
    for x1, y1, x2, y2, title, items, fill, accent in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=28, fill="#" + fill, outline="#" + accent, width=4)
        centered(d, (x1, y1 + 24, x2, y1 + 95), title, pil_font(31, True), "#" + accent)
        for i, item in enumerate(items):
            d.text((x1 + 62, y1 + 125 + i * 52), "✓ " + item, font=pil_font(21), fill="#" + INK)
    arrow(d, (510, 340), (680, 340))
    arrow(d, (1120, 340), (1290, 340))
    arrow(d, (1485, 500), (900, 590), color="#B87800")
    arrow(d, (900, 590), (290, 500), color="#15956A")
    centered(d, (420, 555, 1390, 665), "Session 公共事实互通；私有上下文隔离；每次判断都带预算与证据", pil_font(22, True), "#" + DARK_BLUE)
    im.save(ASSETS / "role-model.png")


def build_flow_diagram():
    im = Image.new("RGB", (1800, 430), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 30), "协作收敛路径", font=pil_font(34, True), fill="#" + INK)
    steps = ["主管分派", "专家并行", "Session Handoff", "裁判核验", "定向返工", "主管汇总"]
    x = 55
    for idx, step in enumerate(steps):
        fill = GREEN_FILL if idx in (1, 5) else BLUE_FILL if idx < 4 else GOLD_FILL
        accent = GREEN if idx in (1, 5) else BRAND_BLUE if idx < 4 else GOLD
        d.rounded_rectangle((x, 150, x + 235, 285), radius=22, fill="#" + fill, outline="#" + accent, width=3)
        centered(d, (x + 8, 160, x + 227, 245), step, pil_font(23, True), "#" + accent)
        centered(d, (x + 8, 235, x + 227, 275), f"阶段 {idx + 1}", pil_font(16), "#" + MUTED)
        if idx < len(steps) - 1:
            arrow(d, (x + 235, 218), (x + 285, 218), color="#66758A", width=4)
        x += 285
    d.text((62, 350), "裁判不修改交付物；证据不足时只向责任 Session 发出明确返工指令，避免全员重跑。", font=pil_font(21), fill="#" + DARK_BLUE)
    im.save(ASSETS / "collaboration-flow.png")


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    set_style_font(title)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb(INK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.05

    subtitle = styles["Subtitle"]
    set_style_font(subtitle)
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        "Heading 1": (16, HEADING_BLUE, 16, 8),
        "Heading 2": (13, HEADING_BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        set_style_font(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    set_style_font(caption)
    caption.font.size = Pt(8.5)
    caption.font.italic = False
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.line_spacing = 1.05
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = False

    if "Cover Kicker" not in styles:
        cover_kicker = styles.add_style("Cover Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cover_kicker = styles["Cover Kicker"]
    set_style_font(cover_kicker)
    cover_kicker.font.size = Pt(10)
    cover_kicker.font.bold = True
    cover_kicker.font.color.rgb = rgb(BRAND_BLUE)
    cover_kicker.paragraph_format.space_after = Pt(18)
    cover_kicker.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Running header/footer (cover remains blank via first-page parts).
    header = section.header
    hp = header.paragraphs[0]
    set_para(hp, after=0, line=1, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_run(hp.add_run("BudgetLoop  |  可治理的专业 Agent 团队控制面"), 8.3, True, MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    set_para(fp, after=0, line=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_run(fp.add_run("第 "), 8.5, False, MUTED)
    add_field(fp, "PAGE")
    set_run(fp.add_run(" 页  /  共 8 页"), 8.5, False, MUTED)

    doc.core_properties.title = "BudgetLoop 可治理的专业 Agent 团队控制面"
    doc.core_properties.subject = "智能路由、Session 协作、预算审计与人机治理客户方案"
    doc.core_properties.author = "BudgetLoop"
    doc.core_properties.keywords = "BudgetLoop, Agent Team, OpenHands, Codex, Gemini CLI, 预算, 审计"
    return doc


def build_docx():
    ASSETS.mkdir(parents=True, exist_ok=True)
    crop_screenshots()
    build_architecture_diagram()
    build_role_diagram()
    build_flow_diagram()
    doc = setup_document()

    # Page 1 — proposal_centerpiece cover.
    add_text(doc, "CUSTOMER & PARTNER BRIEF", 10, True, BRAND_BLUE, before=34, after=30,
             line=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("BudgetLoop")
    p2 = doc.add_paragraph(style="Title")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(12)
    p2.add_run("可治理的专业 Agent 团队控制面")
    sp = doc.add_paragraph(style="Subtitle")
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.add_run("智能选型 · Session 互通 · 专业 Skills · 预算审计 · 人工介入")
    add_callout(
        doc,
        "核心定位",
        "BudgetLoop 不重新打造 Agent 框架。它在 OpenHands、Codex、Gemini CLI 等执行引擎之上，"
        "通过智能路由自动选型，以 Session 互通与专业 Skills 组建团队，并用预算、审计与人工介入保证可控交付。",
        GREEN_FILL,
        GREEN,
        after=18,
    )
    add_fact_strip(doc, [
        ("预算", "可控", "团队与角色两级", BRAND_BLUE),
        ("过程", "可审计", "消息、工具、证据", GREEN),
        ("治理", "可介入", "暂停、纠偏、审批", GOLD),
    ])
    add_text(doc, "主管 Agent 负责编排与汇总；裁判 Agent 独立控制对话预算、核验证据并触发定向返工。",
             11.5, True, DARK_BLUE, before=22, after=24, line=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "面向客户与合作伙伴的目标架构方案", 10, True, MUTED, before=8, after=4,
             line=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "github.com/bei666qi-pan/BudgetLoop  ·  2026", 9, False, LIGHT_MUTED,
             after=0, line=1, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Page 2.
    page_break(doc)
    add_page_title(doc, 2, "产品定位：选择和治理 Agent，而不是替代 Agent",
                   "用户只需描述目标；BudgetLoop 负责选择团队、引擎、预算和协作阶段。")
    add_matrix(doc, ["维度", "传统单 Agent", "松散多 Agent", "BudgetLoop"], [
        ("选型", "用户手动选择", "每个窗口各自决定", "智能路由推荐团队与引擎"),
        ("上下文", "单一长上下文", "复制粘贴易污染", "公共事实互通、私有上下文隔离"),
        ("治理", "执行后看消耗", "缺少统一边界", "调用前预算、审批与审计"),
        ("收敛", "依赖模型自评", "无人统一验收", "裁判判定，主管汇总"),
    ], [1200, 2450, 2450, 3260], font_size=8.8)
    add_image(doc, ASSETS / "smart-routing.png", 6.25,
              "BudgetLoop 真实智能推荐界面：输入季度经营分析目标后推荐商业增长团队与数据分析团队",
              "真实界面｜自然语言目标自动匹配团队；创建前仍可调整角色、预算与协作模式。")
    add_callout(doc, "客户体验", "无需先理解 Agent 框架差异，也无需手工拼接多个聊天窗口；系统先给出可解释推荐，再由用户确认。",
                BLUE_FILL, BRAND_BLUE, after=0)

    # Page 3.
    page_break(doc)
    add_page_title(doc, 3, "目标架构：自动路由、统一 Session、可替换引擎",
                   "执行引擎负责推理与动作；BudgetLoop 保留任务状态、预算、审批、工作区和终态控制权。")
    add_image(doc, ASSETS / "architecture.png", 6.25,
              "BudgetLoop 目标架构图：从用户目标到智能路由、控制平面、专业团队和工具证据回流",
              "架构原则｜引擎可以替换，业务事实与治理边界不能旁路。")
    add_image(doc, ASSETS / "engine-selection.png", 6.25,
              "BudgetLoop 真实执行引擎选择界面，展示 OpenHands、Codex、Gemini CLI 与扩展适配器状态",
              "真实界面｜OpenHands、Codex、Gemini CLI 已进入统一引擎目录；不可用引擎不会静默回退。")

    # Page 4.
    page_break(doc)
    add_page_title(doc, 4, "专业团队：主管、专家与裁判三层协作",
                   "把“谁负责什么、何时交接、谁有权判定完成”写进协作协议。")
    add_image(doc, ASSETS / "role-model.png", 6.25,
              "主管 Agent、专业 Agent 与裁判 Agent 的三层职责模型",
              "职责分离｜主管管全局、专家做交付、裁判保持独立。")
    add_matrix(doc, ["角色", "核心责任", "可见输入", "主要输出"], [
        ("主管 Agent", "拆解、分配、Handoff、冲突处理", "目标、团队状态、裁判结论", "统一计划与最终汇总"),
        ("专业 Agent", "按角色 Skills 执行并自证", "私有上下文、公共事实、工具", "工件、差异、测试与说明"),
        ("裁判 Agent", "核验证据并控制对话预算", "事实证据、验收条件、预算", "通过 / 返工 / 阻塞"),
    ], [1450, 2850, 2750, 2310], font_size=8.9)
    add_callout(doc, "裁判边界",
                "裁判不修改交付物；它独立控制对话 Token、自动回复轮数、消息频率与阶段预算，只向责任 Session 发出定向返工。",
                GOLD_FILL, GOLD, after=0)

    # Page 5.
    page_break(doc)
    add_page_title(doc, 5, "协作闭环：让团队持续收敛，而不是无限对话",
                   "每次 Handoff 都携带公开成果、剩余问题、预算状态和下一步责任。")
    add_image(doc, ASSETS / "collaboration-flow.png", 6.25,
              "主管分派、专家并行、Session Handoff、裁判核验、定向返工和主管汇总的协作流程",
              "闭环路径｜返工只唤醒必要角色，避免全员重复消耗。")
    add_image(doc, ASSETS / "agent-team.png", 6.15,
              "BudgetLoop 真实 Agent Team 界面：六个 Session 状态与自主 Handoff 记录入口",
              "真实界面｜六个 Session 独立运行并保留 Handoff；已裁去运行标识和工作区路径，未展示生成产品成果。")
    add_callout(doc, "人工控制", "操作员可在团队或单个 Session 层级暂停、恢复、纠偏、审批或终止；所有动作进入审计记录。",
                GREEN_FILL, GREEN, after=0)

    # Page 6.
    page_break(doc)
    add_page_title(doc, 6, "治理闭环：预算、审计、介入与可信失败",
                   "预算不是事后仪表盘，而是执行前的控制权；完成必须由可观察证据支撑。")
    add_image(doc, ASSETS / "run-observatory.png", 6.25,
              "BudgetLoop 真实运行指挥台：71.9k/200k Token、9/20 次调用、正常压力模式与 100% 成功率",
              "真实界面｜该 Session 在预算内完成：71.9k / 200k Token、9 / 20 次调用、100% 调用成功率。")
    add_matrix(doc, ["预算维度", "调用前控制", "运行中信号", "超限或异常处理"], [
        ("Token", "预留可用额度", "已用 + 预留 + 速率", "保守策略或停止"),
        ("调用", "检查次数与并发", "成功率、重试、重复动作", "冻结低收益调用"),
        ("费用", "按可用价格预估", "真实结算与未定价标识", "拒绝夸大成本结论"),
        ("时间", "绝对截止 + 活跃时间", "阶段耗时与等待状态", "部分完成并可交接"),
    ], [1450, 2550, 2800, 2560], font_size=8.8)
    add_callout(doc, "可信失败", "预算不足、工具超时或外部依赖不可用时，系统输出已完成内容、未通过验收、证据和建议下一步，不伪装成功。",
                GOLD_FILL, GOLD, after=0)

    # Page 7.
    page_break(doc)
    add_page_title(doc, 7, "多场景能力：办公、数据与软件工程",
                   "同一控制协议复用不同 Skills 与工具链；本页只说明工作方式，不展示生成产品成果。")
    add_matrix(doc, ["场景", "典型交付", "专业团队", "验证证据"], [
        ("办公交付", "PPT、Word、表格、经营分析、调研报告", "研究、分析、编辑、审校、汇总", "页数、结构、公式、引用、渲染"),
        ("数据工作", "清洗、指标分析、图表、仪表盘、复核", "业务分析、数据工程、可视化、复核", "行数、口径、查询、图表与异常"),
        ("软件工程", "需求、架构、开发、测试、修复、发布", "产品、架构、前后端、QA、发布", "测试、构建、Diff、部署与回滚"),
    ], [1300, 3000, 2750, 2310], font_size=9.0)
    add_heading(doc, "专业 Skills 让 Agent 从“会回答”升级为“会交付”", 2)
    add_matrix(doc, ["工作阶段", "Skills 的作用", "BudgetLoop 的治理"], [
        ("理解任务", "读取行业方法、模板和验收规范", "智能路由选择合适团队与引擎"),
        ("执行制作", "调用文档、表格、幻灯片、代码与浏览器工具", "独立工作区、角色预算与审批"),
        ("独立验证", "渲染、测试、数据复核与可访问性检查", "裁判读取事实并判定返工或通过"),
        ("统一交付", "整理工件、摘要、限制和后续建议", "主管汇总跨 Session 工作成果"),
    ], [1550, 3900, 3910], font_size=9.0)
    add_fact_strip(doc, [
        ("办公", "一套协议", "多种交付格式", BRAND_BLUE),
        ("数据", "事实驱动", "口径可追溯", GREEN),
        ("开发", "测试门禁", "失败可交接", GOLD),
    ])
    add_callout(doc, "适用边界", "任务只要能够被工具执行、被可观察结果验证，就可以纳入同一套预算、审计与人工治理闭环。",
                BLUE_FILL, BRAND_BLUE, after=0)

    # Page 8.
    page_break(doc)
    add_page_title(doc, 8, "合作落地：从高频任务开始，固化为专业团队模板",
                   "先在一个可度量场景中验证成本、质量和人工介入，再复制到更多工作流。")
    add_fact_strip(doc, [
        ("选型门槛", "更低", "目标驱动而非框架驱动", BRAND_BLUE),
        ("成本边界", "更清晰", "角色与团队两级预算", GREEN),
        ("过程透明", "可回放", "Session、工具与证据", DARK_BLUE),
        ("业务责任", "仍在人", "审批与随时介入", GOLD),
    ])
    add_heading(doc, "建议试点路径", 2)
    add_matrix(doc, ["阶段", "客户动作", "BudgetLoop 配置", "阶段产出"], [
        ("1 选择", "挑选一个高频、可验收任务", "定义目标、附件与验收条件", "清晰基线"),
        ("2 配置", "确认角色、风险与预算", "智能路由 + 团队模板 + 引擎", "受控团队"),
        ("3 运行", "在审批点介入", "Session 协作 + 裁判闭环 + 审计", "真实证据"),
        ("4 固化", "复盘质量、成本和人工步骤", "沉淀 Skills、预算策略与模板", "可复制流程"),
    ], [1100, 2700, 3200, 2360], font_size=8.9)
    add_heading(doc, "部署与扩展", 2)
    add_matrix(doc, ["能力", "客户价值"], [
        ("自托管与隔离", "数据、凭据和工作区边界由组织掌握；高风险动作保留审批。"),
        ("引擎可替换", "OpenHands、Codex、Gemini CLI 等执行引擎统一接入，不绑定单一供应商。"),
        ("Skills 可扩展", "将行业方法、企业模板和验收标准沉淀为专业团队能力。"),
        ("Session 可审计", "消息、Handoff、工具事实、预算变化和最终报告均可追溯。"),
    ], [2200, 7160], font_size=9.1)
    add_callout(doc, "合作建议", "以 2-4 周试点验证一个办公、数据或软件交付流程；用完成率、总用量、返工次数和人工介入点共同评估。项目仓库与联系：github.com/bei666qi-pan/BudgetLoop",
                GREEN_FILL, GREEN, after=10)
    add_text(doc, "BudgetLoop 让企业使用多个 Agent 时，不再依赖“相信模型”，而是依赖可预算、可审计、可介入的交付系统。",
             11.5, True, DARK_BLUE, before=5, after=18, line=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)

    register_document_font(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_docx()
